import csv
import json
from datetime import UTC, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


def safe_cell(value):
    """Treat metadata as text, never as spreadsheet formulas."""
    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False)
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def export(data, path):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.suffix.lower() == ".json":
        path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
        return
    rows = []
    for work in data.get("publications", []):
        rows.append(
            {
                "id": work["id"],
                "doi": work.get("doi"),
                "title": work["title"],
                "year": work.get("year"),
                "journal": work.get("journal"),
                "issns": work.get("issns", []),
                "authors": [a["name"] for a in work.get("authors", [])],
                "citations_by_source": work.get("citations", []),
                "source_ids": work.get("source_ids", {}),
                "provenance": work.get("provenance", []),
            }
        )
    headers = (
        list(rows[0])
        if rows
        else [
            "id",
            "doi",
            "title",
            "year",
            "journal",
            "issns",
            "authors",
            "citations_by_source",
            "source_ids",
            "provenance",
        ]
    )
    if path.suffix.lower() == ".csv":
        with path.open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.DictWriter(stream, fieldnames=headers)
            writer.writeheader()
            writer.writerows({k: safe_cell(v) for k, v in row.items()} for row in rows)
        export(data, path.with_suffix(".profile.json"))
    elif path.suffix.lower() == ".xlsx":
        book = Workbook()
        sheet = book.active
        sheet.title = "Publications"
        sources = sorted(
            {m["source"] for row in rows for m in row["citations_by_source"] if m["kind"] == "citation_count"}
        )
        sheet.append(
            headers[:7]
            + [item for source in sources for item in (f"{source}: citations", f"{source}: retrieved UTC")]
        )
        for row in rows:
            values = [safe_cell(row.get(k)) for k in headers[:7]]
            for source in sources:
                observations = [
                    m
                    for m in row["citations_by_source"]
                    if m["source"] == source and m["kind"] == "citation_count"
                ]
                try:
                    dated = [(datetime.fromisoformat(m["retrieved_at"]), m) for m in observations]
                    if any(t.tzinfo is None for t, _ in dated):
                        raise ValueError("Naive timestamp")
                    latest = max(t.astimezone(UTC) for t, _ in dated)
                    counts = [m["value"] for t, m in dated if t.astimezone(UTC) == latest]
                    valid = all(isinstance(v, int) and not isinstance(v, bool) and v >= 0 for v in counts)
                    value = counts[0] if valid and len(set(counts)) == 1 else "conflict/invalid"
                    values.extend([value, latest.replace(tzinfo=None)])
                except (ValueError, TypeError):
                    values.extend(["n/a", None])
            sheet.append(values)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = 24
        sheet.column_dimensions["C"].width = 65
        sheet.column_dimensions["D"].width = 9
        sheet.column_dimensions["E"].width = 38
        sheet.column_dimensions["G"].width = 38
        for row in sheet.iter_rows(min_row=2):
            sheet.row_dimensions[row[0].row].height = 66
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if isinstance(cell.value, datetime):
                    cell.number_format = "yyyy-mm-dd hh:mm"
        meta = book.create_sheet("Profile and coverage")
        meta.append(["Profile and corpus", "Value", "Source / meaning"])
        author, coverage = data.get("author", {}), data.get("coverage", {})
        meta.append(["Author", safe_cell(author.get("name", "")), safe_cell(author.get("id", ""))])
        meta.append(["Publications retrieved", len(rows), "Sample; not a universal publication count"])
        meta.append(["Primary source total", coverage.get("reported_total"), coverage.get("source", "")])
        meta.append(
            [
                "Corpus completeness",
                "complete" if coverage.get("complete_in_primary_source") else "partial/unknown",
                "Completeness only within the primary source",
            ]
        )
        meta.append(
            ["Full evidence", path.with_suffix(".profile.json").name, "Original responses and timestamps"]
        )
        meta.append([])
        meta.append(["Metric", "Value", "Source / scope"])
        for metric in author.get("metrics", []):
            meta.append([metric["kind"], safe_cell(metric["value"]), f"{metric['source']}: provider profile"])
        for metric in data.get("calculated_metrics", []):
            meta.append(
                [
                    "h-index of retrieved subset",
                    metric["value"] if metric["value"] is not None else "n/a",
                    f"{metric['source']}: {metric['publications_with_citations']} works with usable counts",
                ]
            )
        meta.append([])
        meta.append(["Provider", "Status", "Details"])
        for provider in data.get("providers", []):
            meta.append([provider["source"], provider["status"], safe_cell(provider.get("reason", ""))])
        meta.column_dimensions["A"].width = 32
        meta.column_dimensions["B"].width = 34
        meta.column_dimensions["C"].width = 70
        meta.freeze_panes = "A2"
        for row in meta.iter_rows():
            meta.row_dimensions[row[0].row].height = 31
            for cell in row:
                cell.alignment = Alignment(vertical="center", wrap_text=True)
        for ws in (sheet, meta):
            ws.sheet_view.showGridLines = False
            for cell in ws[1]:
                cell.fill = PatternFill("solid", fgColor="243C56")
                cell.font = Font(color="FFFFFF", bold=True)
                cell.alignment = Alignment(wrap_text=True, vertical="center")
            ws.row_dimensions[1].height = 32
        book.save(path)
        export(data, path.with_suffix(".profile.json"))
    else:
        raise ValueError("Export extension must be .json, .csv or .xlsx")
